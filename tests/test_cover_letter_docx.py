from io import BytesIO

import pytest
from docx import Document

from job_agent.services.cover_letter_docx import CoverLetterExportError, export_cover_letter_docx


def test_export_preserves_template_page_setup_and_replaces_only_body() -> None:
    letter = (
        "First tailored paragraph.\n\nSecond tailored paragraph.\n\nThird tailored paragraph."
        "\n\nFourth tailored paragraph.\n\nFifth tailored paragraph."
    )
    exported = Document(BytesIO(export_cover_letter_docx(letter)))
    section = exported.sections[0]

    assert section.top_margin.inches == 1
    assert section.left_margin.inches == 1.25
    assert [paragraph.text for paragraph in exported.paragraphs if "tailored paragraph" in paragraph.text] == [
        "First tailored paragraph.",
        "Second tailored paragraph.",
        "Third tailored paragraph.",
        "Fourth tailored paragraph.",
        "Fifth tailored paragraph.",
    ]


def test_export_rejects_a_cover_letter_without_template_body_paragraphs() -> None:
    with pytest.raises(CoverLetterExportError, match="three to five"):
        export_cover_letter_docx("Only one paragraph.")


def test_export_does_not_duplicate_a_legacy_salutation_or_closing() -> None:
    letter = (
        "Dear Hiring Manager,\n\nFirst paragraph.\n\nSecond paragraph.\n\nThird paragraph."
        "\n\nKind regards,\n\nCandidate Name"
    )
    exported = Document(BytesIO(export_cover_letter_docx(letter)))
    rendered = [paragraph.text for paragraph in exported.paragraphs]

    assert rendered.count("Dear Hiring Manager,") == 1
    assert rendered.count("Kind regards,") == 1


def test_export_uses_supplied_salutation_and_signature_with_template_layout() -> None:
    letter = (
        "Dear Hiring Team,\n\nFirst paragraph.\n\nSecond paragraph.\n\nThird paragraph."
        "\n\nKind regards,\n\nAlex Candidate"
    )
    exported = Document(BytesIO(export_cover_letter_docx(letter)))
    rendered = [paragraph.text for paragraph in exported.paragraphs]

    assert "Dear Hiring Team," in rendered
    assert "Alex Candidate" in rendered
